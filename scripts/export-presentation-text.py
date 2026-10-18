from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/NOI-DUNG-THUYET-TRINH-BUOC-1-2.md'
OUTPUT = SOURCE.with_suffix('.docx')


def plain(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
    return text.replace('**', '').replace('`', '')


def inline(paragraph, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        else:
            paragraph.add_run(token)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(1.6)
section.right_margin = Cm(1.6)
section.header_distance = Cm(.7)
section.footer_distance = Cm(.7)

for name in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 24), ('Heading 1', 17), ('Heading 2', 14), ('Heading 3', 12)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string('17365D')

header = section.header.paragraphs[0]
header.text = 'QUẢN LÝ NHÀ TRỌ | PHÂN TÍCH YÊU CẦU & THIẾT KẾ CSDL'
header.runs[0].font.name = 'Arial'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('64748B')
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Trang ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

doc.core_properties.title = 'Nội dung thuyết trình: Phân tích yêu cầu và thiết kế CSDL quản lý nhà trọ'
doc.core_properties.subject = 'Nội dung 20 slide chính và hai trang ERD phụ lục'
doc.core_properties.author = 'Nhóm thực hiện'

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line or line == '---':
        i += 1
        continue

    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            row = [cell.strip() for cell in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', cell) for cell in row):
                rows.append(row)
            i += 1
        columns = max(map(len, rows))
        table = doc.add_table(rows=1, cols=columns)
        table.style = 'Light Shading Accent 1'
        for index, row in enumerate(rows):
            cells = table.rows[0].cells if index == 0 else table.add_row().cells
            for cell_index, text in enumerate(row):
                paragraph = cells[cell_index].paragraphs[0]
                inline(paragraph, text)
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if index == 0:
                        run.bold = True
            properties = table.rows[index]._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            if index == 0:
                properties.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
        continue

    heading = re.match(r'^(#{1,3})\s+(.*)', line)
    if heading:
        hashes, text = heading.groups()
        if text.startswith('Slide ') or text.startswith('Phụ lục '):
            doc.add_page_break()
        if i == 0:
            doc.add_paragraph(text, style='Title')
        else:
            doc.add_heading(text, level=min(len(hashes), 3))
        i += 1
        continue

    picture = re.match(r'Chèn \[[^\]]+\]\(([^)]+)\)', line)
    if picture:
        image_path = SOURCE.parent / picture.group(1)
        if image_path.suffix == '.svg':
            image_path = image_path.with_suffix('.png')
        doc.add_picture(str(image_path), width=Cm(17.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_path.name == 'quanlytro-dbdiagram.png':
            doc.add_paragraph('ERD đầy đủ; dùng ảnh gốc hoặc PDF khi phóng to từng cụm bảng.', style='Caption')
        i += 1
        continue

    if line.startswith('- '):
        paragraph = doc.add_paragraph(style='List Bullet')
        inline(paragraph, line[2:])
    elif re.match(r'^\d+\. ', line):
        paragraph = doc.add_paragraph(style='List Number')
        inline(paragraph, re.sub(r'^\d+\. ', '', line))
    else:
        paragraph = doc.add_paragraph()
        inline(paragraph, line)
    i += 1

doc.save(OUTPUT)
print(OUTPUT.relative_to(ROOT))
print(f'{len(doc.paragraphs)} paragraphs; {len(doc.tables)} tables; {len(doc.inline_shapes)} embedded diagrams')
