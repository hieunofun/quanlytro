"""Generate the reviewable Word report and PowerPoint slides from local project evidence."""
from pathlib import Path
import re
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.util import Inches, Pt as SlidePt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'

def plain(text):
    text = re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', text)
    return text.replace('**', '').replace('`', '')

def report():
    doc = Document()
    doc.sections[0].top_margin = Cm(2)
    doc.sections[0].bottom_margin = Cm(2)
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    doc.core_properties.title = 'Hệ thống quản lý nhà trọ - Báo cáo đồ án Java/Web'
    for filename in ['REPORT.md', 'DEMO.md', 'V1-COVERAGE.md', 'VALIDATION.md']:
        if filename != 'REPORT.md':
            doc.add_page_break()
        lines = (DOCS / filename).read_text(encoding='utf-8').splitlines()
        code = False
        diagram_index = 0
        table_lines = []
        def flush_table():
            if not table_lines:
                return
            data = [[plain(c.strip()) for c in row.strip('|').split('|')] for row in table_lines if not re.match(r'^\|[\s:|-]+\|$', row)]
            table = doc.add_table(rows=0, cols=max(map(len, data)))
            table.style = 'Light Shading Accent 1'
            for row in data:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = value
            table_lines.clear()
        for line in lines:
            if line.startswith('```'):
                flush_table()
                code = not code
                if code:
                    names=['architecture','data-model']
                    if filename=='REPORT.md' and diagram_index<len(names) and (DOCS/'diagrams'/(names[diagram_index]+'.png')).exists():
                        doc.add_picture(str(DOCS/'diagrams'/(names[diagram_index]+'.png')),width=Cm(15))
                        diagram_index+=1
                    else:
                        doc.add_paragraph('Mã nguồn định dạng xem trong file Markdown tương ứng.', style='Caption')
                continue
            if code:
                continue
            if line.startswith('|'):
                table_lines.append(line)
                continue
            flush_table()
            if not line.strip():
                continue
            if line.startswith('# '):
                doc.add_heading(plain(line[2:]), 0)
            elif line.startswith('## '):
                doc.add_heading(plain(line[3:]), 1)
            elif line.startswith('- '):
                doc.add_paragraph(plain(line[2:]), style='List Bullet')
            else:
                doc.add_paragraph(plain(line))
        flush_table()
    for name, caption in [('dashboard-desktop.png', 'Dashboard quản lý - dữ liệu demo thật từ API'), ('invoice-desktop.png', 'Hóa đơn với từng công thức và lịch sử thanh toán'), ('tenant-mobile.png', 'Cổng người thuê tại chiều rộng 375 px')]:
        image = DOCS / 'screenshots' / name
        if image.exists():
            doc.add_page_break()
            doc.add_heading(caption, 1)
            with Image.open(image) as source:
                width = min(15, 22 * source.width / source.height)
            doc.add_picture(str(image), width=Cm(width))
    doc.save(DOCS / 'bao-cao-do-an.docx')

def slides():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    content = [
        ('QUẢN LÝ NHÀ TRỌ', ['Hóa đơn minh bạch và cảnh báo tiêu thụ điện nước bất thường', 'Đồ án Java/Web · Spring Boot / React / PostgreSQL', 'Điền họ tên, MSSV, lớp và giảng viên trước khi nộp']),
        ('Bài toán và người dùng', ['Chủ nhà: quản lý 1–3 khu, theo dõi phòng và công nợ', 'Nhân viên: chỉ truy cập khu hoặc phòng cụ thể được phân công', 'Người thuê: xem hợp đồng, công thức tính tiền, báo sự cố']),
        ('Kiến trúc', ['React → REST API → Spring Boot → PostgreSQL', 'Controller / DTO → Service → JPA Repository + JDBC', 'JWT và Access kiểm tra vai trò, phạm vi ở backend', 'Flyway quản lý schema; AOP đo thời gian; audit cùng transaction']),
        ('Luồng nghiệp vụ', ['Phòng → người thuê → hợp đồng và tiền cọc', 'Ghi chỉ số → kiểm tra cảnh báo → phát hành hóa đơn', 'Thanh toán nhiều lần → công nợ → trả phòng và cấn trừ cọc']),
        ('Hóa đơn minh bạch', ['Nước: (128 − 120) × 20.000 = 160.000 VND', 'Lưu snapshot đơn giá, chỉ số, lượng dùng và công thức', 'Tạo giá mới không làm thay đổi hóa đơn đã phát hành', 'Tiền phòng/dịch vụ chia theo ngày; dịch vụ đầu người dùng ngày-người']),
        ('Cảnh báo và sửa chỉ số', ['Lịch sử: 8, 9, 10 m³ → trung bình 9 m³', 'Kỳ mới: 27 m³ → tăng 200%, vượt ngưỡng 50%', 'Sửa 174 → 155: tiêu thụ còn 8 m³, lưu lịch sử sửa', 'Chặn phát hành khi cảnh báo còn chờ kiểm tra']),
        ('Transaction và bảo mật', ['Unique hóa đơn/hợp đồng/kỳ, exclusion constraint hợp đồng', 'Khóa hàng khi thu tiền; không thu vượt công nợ', 'JWT 30 phút, refresh xoay; logout/đổi mật khẩu thu hồi phiên', 'Sai 5 lần: khóa 15 phút; reset mã một lần do chủ nhà cấp']),
        ('Demo trên dữ liệu thật', ['Hóa đơn 09/2026: 3.690.000 VND', 'Ghi nhận 2.000.000 → còn nợ 1.690.000 VND', 'Đăng nhập người thuê để xem chi tiết', 'Thử hóa đơn trùng, đổi giá kỳ sau, kiểm tra quyền nhân viên']),
        ('Kiểm thử và bàn giao', ['JUnit + SpringBootTest trên PostgreSQL thật', 'HTTP tests: JWT, refresh, logout, khóa tài khoản, đổi mật khẩu', 'PDF tiếng Việt; trình duyệt desktop và 375 px', 'JAR, source ZIP, script demo, API, báo cáo và cấu hình Docker']),
        ('Hoàn thiện V1 và hướng phát triển', ['Ảnh tải từ máy có kiểm tra và phân quyền; vi phạm có workflow', 'Phân công theo khu/phòng; áp dụng lịch chốt và phát hành', 'Báo cáo tiền thu theo ngày nhận, tách cấn trừ cọc; kiểm chứng Docker', 'Mobile/ngân hàng/AI/IoT ngoài V1; mở rộng phân trang và phiếu điều chỉnh']),
    ]
    for i, (title, bullets) in enumerate(content, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(247, 248, 250)
        title_box = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(12), Inches(1)).text_frame
        title_box.text = title
        for p in title_box.paragraphs:
            p.font.name = 'Arial'
            p.font.size = SlidePt(30)
            p.font.bold = True
            p.font.color.rgb = RGBColor(24, 24, 24)
        body = slide.shapes.add_textbox(Inches(.8), Inches(1.8), Inches(11.7), Inches(4.6)).text_frame
        body.word_wrap = True
        for j, bullet in enumerate(bullets):
            p = body.paragraphs[0] if j == 0 else body.add_paragraph()
            p.text = '• ' + bullet
            p.font.name = 'Arial'
            p.font.size = SlidePt(23)
            p.font.color.rgb = RGBColor(44, 44, 44)
            p.space_after = SlidePt(24)
        footer = slide.shapes.add_textbox(Inches(.8), Inches(6.9), Inches(12), Inches(.3)).text_frame
        footer.text = f'NHÀ TRỌ · ĐỒ ÁN JAVA/WEB                                                        {i:02d} / {len(content):02d}'
        footer.paragraphs[0].font.size = SlidePt(11)
        footer.paragraphs[0].font.color.rgb = RGBColor(112, 112, 112)
    prs.save(DOCS / 'slide-bao-ve.pptx')

if __name__ == '__main__':
    report()
    slides()
    print('Created docs/bao-cao-do-an.docx and docs/slide-bao-ve.pptx')
