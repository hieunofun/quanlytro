"""Build the editable classroom deck and its matching, note-free content document."""
import json
from pathlib import Path
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
OUT=DOCS/'thiet-ke-csdl-va-phan-tich-nghiep-vu.pptx'
INK='17212B'; ACCENT='284A63'; MUTED='59636E'; BORDER='C9D0D6'; PALE='F0F3F5'; WHITE='FFFFFF'
W,H=13.333333,7.5
prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
prs.core_properties.title='Hệ thống quản lý nhà trọ — Phân tích yêu cầu và thiết kế CSDL'
prs.core_properties.subject='Thiết kế mục tiêu 17 bảng — hồ sơ bước 1–2'
prs.core_properties.author='Nhóm thực hiện'
prs.core_properties.keywords='quản lý nhà trọ, yêu cầu nghiệp vụ, ERD, chuẩn hóa'
CONTENT=[]


def no_effects(o):
    o._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in o._element.xpath('./p:style/a:effectRef'):effect.set('idx','0')


def rect(s,x,y,w,h,fill=WHITE,line=BORDER,kind=MSO_SHAPE.RECTANGLE):
    o=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    o.fill.solid();o.fill.fore_color.rgb=RGBColor.from_string(fill)
    if line: o.line.color.rgb=RGBColor.from_string(line);o.line.width=Pt(1)
    else:o.line.fill.background()
    no_effects(o)
    return o


def text(s,x,y,w,h,value,size=22,bold=False,color=INK,align=PP_ALIGN.LEFT,font='Arial'):
    o=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=o.text_frame
    tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for n,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if n==0 else tf.add_paragraph()
        p.text=line;p.alignment=align;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(5 if n<len(value.split('\n'))-1 else 0)
        p.line_spacing=1.08
    return o


def line(s,x1,y1,x2,y2,arrow=False):
    o=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    o.line.color.rgb=RGBColor.from_string(MUTED);o.line.width=Pt(1.15)
    no_effects(o)
    if arrow:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');o.line._get_or_add_ln().append(end)
    return o


def route(s,points,label=None,label_at=None):
    for i in range(len(points)-1):line(s,*points[i],*points[i+1],arrow=i==len(points)-2)
    if label:
        x,y,w=label_at
        rect(s,x,y,w,.26,WHITE,None)
        text(s,x,y,w,.25,label,12.5,True,MUTED,PP_ALIGN.CENTER)


def base(title,section,items=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    text(s,.62,.25,12,.25,section,10.5,True,ACCENT)
    text(s,.62,.72,12,.52,title,30,True)
    line(s,.62,1.36,12.7,1.36)
    line(s,.62,7.03,12.7,7.03)
    text(s,.62,7.15,11,.18,'QUẢN LÝ NHÀ TRỌ  |  PHÂN TÍCH YÊU CẦU VÀ THIẾT KẾ CSDL',9,False,MUTED)
    text(s,12.1,7.12,.6,.25,f'{len(prs.slides):02}',10,False,MUTED,PP_ALIGN.RIGHT)
    CONTENT.append({'slide':len(prs.slides),'title':title,'section':section,'items':items or []})
    return s


def bullets(s,items,y=1.8,size=22,gap=.8):
    for i,item in enumerate(items):text(s,.8,y+i*gap,11.8,gap-.10,'• '+item,size)


def table(s,rows,widths,y=1.75,row_h=.72,size=18):
    o=s.shapes.add_table(len(rows),len(widths),Inches(.65),Inches(y),Inches(sum(widths)),Inches(row_h*len(rows)))
    t=o.table
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for r,values in enumerate(rows):
        t.rows[r].height=Inches(row_h)
        for j,value in enumerate(values):
            cell=t.cell(r,j);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(ACCENT if r==0 else WHITE)
            cell.margin_left=cell.margin_right=Inches(.13);cell.margin_top=cell.margin_bottom=Inches(.06)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            tf=cell.text_frame;tf.clear();tf.word_wrap=True
            for k,part in enumerate(str(value).split('\n')):
                p=tf.paragraphs[0] if k==0 else tf.add_paragraph();p.text=part
                p.font.name='Arial';p.font.size=Pt(size);p.font.bold=r==0
                p.font.color.rgb=RGBColor.from_string(WHITE if r==0 else INK);p.space_after=Pt(1);p.line_spacing=1.02
    CONTENT[-1]['table']=rows
    return o


def node(s,x,y,w,h,label,decision=False,size=16):
    rect(s,x,y,w,h,PALE if decision else WHITE,ACCENT,MSO_SHAPE.DIAMOND if decision else MSO_SHAPE.RECTANGLE)
    pad=.24 if decision else .10
    o=text(s,x+pad,y+.05,w-2*pad,h-.10,label,size,False,INK,PP_ALIGN.CENTER)
    o.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE


def entity(s,x,y,name,fields,w=3.1):
    h=.50+.35*len(fields)
    rect(s,x,y,w,h)
    rect(s,x,y,w,.38,PALE,BORDER)
    text(s,x+.1,y+.06,w-.2,.26,name,15.5 if len(name)>20 else 17,True)
    for i,f in enumerate(fields):text(s,x+.1,y+.48+i*.35,w-.2,.30,f,14)
    return (x,y,w,h)


def relation(s,a,b,label='1 : N',vertical=False):
    x,y,w,h=a; xx,yy,ww,hh=b
    if vertical:
        if y<yy:route(s,[(x+w/2,y+h),(x+w/2,yy-.2),(xx+ww/2,yy-.2),(xx+ww/2,yy)],label,(x+w/2+.1,(y+h+yy)/2-.15,1.0))
        else:route(s,[(x+w/2,y),(x+w/2,yy+hh+.2),(xx+ww/2,yy+hh+.2),(xx+ww/2,yy+hh)],label,(x+w/2+.1,(y+yy+hh)/2-.15,1.0))
    elif x<xx:route(s,[(x+w,y+.7),(xx,y+.7)],label,((x+w+xx)/2-.5,y+.37,1.0))
    else:route(s,[(x,y+.7),(xx+ww,y+.7)],label,((x+xx+ww)/2-.5,y+.37,1.0))


# 01
s=base('Hệ thống quản lý nhà trọ','ĐỒ ÁN HỌC PHẦN',[
 'Phân tích yêu cầu và thiết kế cơ sở dữ liệu',
 'Hóa đơn minh bạch — theo dõi công nợ — cảnh báo điện nước',
 'Một chủ nhà, nhiều khu; ba tác nhân ADMIN / STAFF / TENANT'])
text(s,.8,2.1,11.6,1.35,'PHÂN TÍCH YÊU CẦU\nVÀ THIẾT KẾ CƠ SỞ DỮ LIỆU',34,True,ACCENT)
text(s,.8,3.65,11.6,.7,'Hợp đồng • Hóa đơn • Công nợ • Quyết toán',24)
text(s,.8,5.0,11.6,.9,'Kỹ năng lập trình nâng cao — N03\nNhóm: ………………………………',19,False,MUTED)

# 02
s=base('Bài toán và phạm vi','01  PHÂN TÍCH YÊU CẦU')
items=['Một chủ nhà quản lý nhiều khu và nhiều phòng.',
 'Theo dõi người thuê, hợp đồng, chỉ số và các khoản phải thu.',
 'Người thuê xem được căn cứ tính tiền và lịch sử thanh toán.',
 'MVP: cho thuê → tính tiền → thu tiền → quyết toán; kèm xử lý sự cố.',
 'Thanh toán ghi nhận thủ công; đặt phòng công khai, AI/IoT ngoài phạm vi.']
bullets(s,items,size=22,gap=.9);CONTENT[-1]['items']=items

# 03
s=base('Tác nhân và phân quyền','01  PHÂN TÍCH YÊU CẦU')
table(s,[['Tác nhân','Được thực hiện','Giới hạn'],
 ['Chủ nhà\nADMIN','Quản lý tài khoản, phân công, khu/phòng và tài chính','Toàn bộ hệ thống của chủ nhà'],
 ['Nhân viên\nSTAFF','Lập hợp đồng, ghi chỉ số, lập hóa đơn, thu tiền, xử lý sự cố','Khu được phân công; không quản trị tài khoản/quyền'],
 ['Người thuê\nTENANT','Xem hợp đồng, hóa đơn, giao dịch; gửi và theo dõi sự cố','Chỉ hợp đồng mình đại diện; không tự xác nhận thu tiền']],
 [2.0,5.7,4.3],row_h=1.12,size=19)
text(s,.8,6.48,11.7,.3,'Thành viên ở cùng là hồ sơ cư trú; không mặc nhiên có quyền của người đại diện.',16,False,MUTED)

# 04
s=base('Chức năng quản lý và cho thuê','01  PHÂN TÍCH YÊU CẦU')
table(s,[['Nhóm chức năng','Use Case chính','Tác nhân'],
 ['Tài khoản','Tạo tài khoản; cập nhật, khóa tài khoản','ADMIN'],
 ['Phân công','Phân công nhân viên; thu hồi phân công','ADMIN'],
 ['Khu và phòng','Tạo khu/phòng; cập nhật phòng; chuyển phòng bảo trì','ADMIN; STAFF theo quyền'],
 ['Người thuê','Tạo, cập nhật hồ sơ; tra cứu người thuê','ADMIN; tra cứu có giới hạn'],
 ['Hợp đồng','Lập, tra cứu hợp đồng; cập nhật thành viên cư trú','ADMIN / STAFF; TENANT tra cứu']],
 [2.1,7.0,2.9],row_h=.77,size=17.5)

# 05
s=base('Chức năng tính tiền và vận hành','01  PHÂN TÍCH YÊU CẦU')
table(s,[['Nhóm chức năng','Use Case chính'],
 ['Dịch vụ','Thiết lập đơn giá và ngưỡng cảnh báo theo khu'],
 ['Điện nước','Ghi chỉ số; sửa chỉ số cuối; kiểm tra cảnh báo'],
 ['Hóa đơn','Lập nháp; xóa nháp; phát hành; tra cứu hóa đơn'],
 ['Thanh toán','Ghi nhận thanh toán; tra cứu công nợ; tổng hợp tiền thu'],
 ['Trả phòng','Chốt kỳ cuối; quyết toán; cấn cọc; xác định hoàn/thu thêm'],
 ['Sự cố','Gửi yêu cầu; đính kèm ảnh; cập nhật và theo dõi tiến độ']],
 [2.2,9.8],row_h=.66,size=18)

# 06: native editable flowchart
s=base('Luồng lập và phát hành hóa đơn','01  PHÂN TÍCH YÊU CẦU',[
 'Kiểm tra kỳ, đơn giá, chỉ số và hóa đơn trùng trước khi lập nháp.',
 'Có cảnh báo chờ kiểm tra thì chưa phát hành; nhập sai thì xóa nháp, sửa chỉ số và lập lại.',
 'Chỉ phát hành khi đến ngày; giữ nguyên giá snapshot và chỉ số được tham chiếu.'])
node(s,.6,2.1,2.4,.8,'Chọn hợp đồng\nvà kỳ')
node(s,3.5,1.9,2.8,1.2,'Đủ dữ liệu,\nchưa trùng kỳ?',True)
node(s,6.8,2.1,2.6,.8,'Tính tiền;\nlưu hóa đơn nháp')
node(s,10.0,1.9,2.8,1.2,'Còn cảnh báo\nPENDING?',True)
route(s,[(3,2.5),(3.5,2.5)]);route(s,[(6.3,2.5),(6.8,2.5)],'Có',(6.35,2.15,.4));route(s,[(9.4,2.5),(10,2.5)])
node(s,3.5,3.75,2.8,.9,'Dừng; bổ sung\nhoặc sửa dữ liệu')
route(s,[(4.9,3.1),(4.9,3.75)],'Không',(5,3.35,.65))
node(s,10,3.75,2.8,1.05,'Quản lý kiểm tra\nXác nhận / sửa chỉ số')
route(s,[(11.4,3.1),(11.4,3.75)],'Có',(11.55,3.35,.4))
node(s,6.8,5.25,2.8,1.15,'Đã đến ngày\nphát hành?',True)
node(s,3.5,5.45,2.8,.8,'Giữ nháp;\nchờ ngày phát hành')
node(s,10,5.45,2.8,.8,'Phát hành ISSUED\nNgười thuê được xem')
route(s,[(11.4,1.9),(11.4,1.65),(13,1.65),(13,5.05),(8.2,5.05),(8.2,5.25)],'Không',(11.7,1.39,.75))
route(s,[(10,4.15),(8.1,4.15),(8.1,2.9)],'Sai: lập lại',(8.35,3.8,1.1))
route(s,[(11.4,4.8),(11.4,5.05),(8.2,5.05),(8.2,5.25)],'Đúng',(11.55,4.85,.65))
route(s,[(6.8,5.82),(6.3,5.82)],'Chưa',(6.3,5.5,.5));route(s,[(9.6,5.82),(10,5.82)],'Rồi',(9.6,5.5,.4))
text(s,.8,6.65,11.9,.25,'Không đủ lịch sử hoặc trung bình = 0: không đánh giá tăng theo phần trăm.',15,False,MUTED)

# 07
s=base('Luồng quyết toán trả phòng','01  PHÂN TÍCH YÊU CẦU',[
 'Kiểm tra quyền, ngày trả và việc quyết toán trùng.',
 'Chốt chỉ số, hoàn tất hóa đơn mọi tháng đã ở; hóa đơn đã phát hành phải khớp ngày/chỉ số.',
 'Số quyết toán = cọc − thiệt hại − công nợ; công nợ đã gồm điện nước kỳ cuối.',
 'Lưu kết quả, tạo giao dịch cấn cọc và kết thúc hợp đồng trong một giao dịch.'])
node(s,.6,2.1,2.4,.8,'Chọn hợp đồng\nvà ngày trả')
node(s,3.5,1.9,2.8,1.2,'Có quyền, hợp lệ,\nchưa quyết toán?',True,15.5)
node(s,6.8,2.1,2.6,.8,'Chốt chỉ số;\nhoàn tất các kỳ')
node(s,10,1.9,2.8,1.2,'Đủ hóa đơn,\nkhớp dữ liệu?',True)
route(s,[(3,2.5),(3.5,2.5)]);route(s,[(6.3,2.5),(6.8,2.5)],'Có',(6.3,2.15,.5));route(s,[(9.4,2.5),(10,2.5)])
node(s,3.5,3.65,2.8,.8,'Dừng; thông báo\nlý do cụ thể')
node(s,10,3.65,2.8,.8,'Dừng; bổ sung\nhoặc kiểm tra lại')
route(s,[(4.9,3.1),(4.9,3.65)],'Không',(5.0,3.3,.65));route(s,[(11.4,3.1),(11.4,3.65)],'Không',(11.55,3.3,.65))
node(s,6.65,4.75,3.1,.75,'Cọc − Thiệt hại − Nợ',size=18)
route(s,[(12.8,2.5),(13,2.5),(13,5.12),(9.75,5.12)],'Có',(12.2,4.85,.5))
node(s,.65,6.0,3.7,.64,'Dương: hoàn tiền',size=18)
node(s,4.8,6.0,3.7,.64,'Bằng 0: không hoàn / thu thêm',size=16)
node(s,8.95,6.0,3.7,.64,'Âm: ghi nhận phải thu thêm',size=17)
route(s,[(8.2,5.5),(8.2,5.72),(2.5,5.72),(2.5,6)])
route(s,[(8.2,5.72),(6.65,5.72),(6.65,6)])
route(s,[(8.2,5.72),(10.8,5.72),(10.8,6)])
text(s,.8,6.75,11.8,.22,'Cấn cọc không phải tiền thu mới. Phần nợ chưa được thanh toán vẫn được giữ lại.',14.5,False,MUTED)

# 08
items=['Khoảng cư trú [ngày vào, ngày ra): trả ngày 16 thì tính đến hết ngày 15.',
 'Phí tháng tính theo ngày ở/ngày-người; điện nước tính chênh chỉ số.',
 'Đơn giá lấy tại đầu kỳ; một hợp đồng có tối đa một hóa đơn mỗi kỳ.',
 'Hóa đơn đã phát hành giữ nguyên; mỗi lần thu phải dương và không vượt nợ.',
 'Cọc trừ thiệt hại, sau đó cấn công nợ; không trừ điện nước hai lần.']
s=base('Quy tắc nghiệp vụ chính','01  PHÂN TÍCH YÊU CẦU',items);bullets(s,items,size=21,gap=.92)

# 09
s=base('Mô hình dữ liệu: 17 bảng','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Nhóm dữ liệu','Các bảng'],
 ['Tài khoản và phân công','users, employee_assignments'],
 ['Khu và phòng','properties, rooms'],
 ['Người thuê và hợp đồng','tenant_profiles, contracts, contract_members'],
 ['Dịch vụ và điện nước','service_items, service_price_configs, meter_readings, consumption_alerts'],
 ['Hóa đơn và quyết toán','invoices, invoice_items, payment_transactions, checkout_records'],
 ['Sự cố và ảnh','maintenance_requests, maintenance_images']],
 [3.0,9.0],row_h=.65,size=17)
text(s,.8,6.55,11.8,.35,'138 cột • 17 PK • 28 FK • 55 CHECK • 2 EXCLUDE • 6 view tính toán',17,True,ACCENT)

# 10
s=base('ERD — khu nhà, người thuê và hợp đồng','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
a=entity(s,.65,1.8,'properties',['PK  id','UQ  property_code'])
b=entity(s,4.95,1.8,'rooms',['PK  id','FK  property_id','room_code; max_occupants'])
c=entity(s,9.25,1.8,'contracts',['PK  id','FK  room_id','FK  primary_tenant_id','start_date; end_date'])
d=entity(s,9.25,4.55,'tenant_profiles',['PK  id','FK/UQ  user_id (tùy chọn)','full_name; identity_number'])
e=entity(s,4.95,4.55,'contract_members',['PK  id','FK  contract_id','FK  tenant_id','effective_from; effective_to'])
relation(s,a,b);relation(s,b,c);relation(s,d,c,vertical=True);relation(s,d,e)
route(s,[(10.4,c[1]+c[3]),(10.4,4.12),(6.5,4.12),(6.5,4.55)],'1 : N',(7.8,3.83,1.0))
text(s,.8,6.65,11.8,.25,'Một người dùng chung hồ sơ qua nhiều lần cư trú; không sao chép họ tên vào bảng thành viên.',16,False,MUTED)
CONTENT[-1]['items']=['properties 1–N rooms; rooms 1–N contracts; tenant_profiles 1–N contracts.',
 'contracts N–N tenant_profiles qua contract_members đối với các lần ở cùng.',
 'Mã phòng duy nhất trong khu; khoảng thuê không được chồng nhau.']

# 11
s=base('ERD — đơn giá, chỉ số và cảnh báo','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
a=entity(s,.65,1.8,'properties',['PK  id','ngưỡng; cửa sổ lịch sử'])
b=entity(s,4.95,1.8,'service_price_configs',['PK  id','FK  property_id','FK  service_item_id','effective_from; unit_price'])
c=entity(s,9.25,1.8,'service_items',['PK  id','UQ  service_code','calculation_type; unit_name'])
d=entity(s,.65,4.5,'contracts',['PK  id','FK  room_id'])
e=entity(s,4.95,4.5,'meter_readings',['PK  id','FK  contract_id; service_item_id','billing_period','old_index; new_index'])
f=entity(s,9.25,4.5,'consumption_alerts',['PK  id','FK/UQ  meter_reading_id','baseline; ngưỡng áp dụng','status; checked_by'])
relation(s,a,b);relation(s,c,b);relation(s,d,e);relation(s,e,f,'1 : 0..1')
route(s,[(10.8,c[1]+c[3]),(10.8,4.05),(6.5,4.05),(6.5,4.5)],'1 : N',(8.2,3.78,1.0))
CONTENT[-1]['items']=['Đơn giá riêng cho từng khu–dịch vụ–ngày hiệu lực.',
 'Chỉ số chỉ lưu hợp đồng, kỳ và dịch vụ đo đếm; phòng và lượng dùng truy ra bằng view.',
 'Một chỉ số có tối đa một cảnh báo; thông tin đánh giá và người kiểm tra có căn cứ.']

# 12
s=base('ERD — hóa đơn, giao dịch và quyết toán','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
a=entity(s,.65,1.8,'contracts',['PK  id','giá thuê; cọc thỏa thuận'])
b=entity(s,4.95,1.8,'invoices',['PK  id','FK  contract_id','billing_period; status'])
c=entity(s,9.25,1.8,'invoice_items',['PK  id','FK  invoice_id','FK  service_item_id / reading','giá snapshot; ngày tính phí'])
d=entity(s,.65,4.5,'checkout_records',['PK  id','FK/UQ  contract_id','cọc / nợ snapshot; thiệt hại'])
e=entity(s,4.95,4.5,'payment_transactions',['PK  id','FK  invoice_id','amount; payment_method'])
f=entity(s,9.25,4.5,'meter_readings',['PK  id','chỉ số đã được khóa'])
relation(s,a,b);relation(s,b,c);relation(s,a,d,'1 : 0..1',True);relation(s,b,e,vertical=True);relation(s,f,c,'0..1',True)
text(s,.8,6.55,11.8,.35,'Tổng tiền, đã thu, công nợ và số quyết toán được tính qua view từ dữ liệu gốc.',17,False,MUTED)
CONTENT[-1]['items']=['Một hóa đơn có nhiều dòng và nhiều giao dịch.',
 'Mỗi hợp đồng có tối đa một quyết toán; dòng đo đếm tham chiếu chỉ số cùng hợp đồng/kỳ.',
 'Chỉ lưu nguồn dữ liệu và snapshot cần bảo toàn; không lưu nhiều bản sao của tổng tiền.']

# 13
s=base('Thuộc tính và kiểu dữ liệu','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Dữ liệu','Kiểu lựa chọn','Căn cứ'],
 ['Khóa nội bộ','BIGINT IDENTITY','Liên kết ổn định bằng PK/FK'],
 ['Tiền, đơn giá','NUMERIC(12,2)','Thập phân chính xác; làm tròn thành tiền'],
 ['Chỉ số điện nước','NUMERIC(10,2)','Có thể có phần thập phân; mới ≥ cũ'],
 ['Điện thoại, giấy tờ','VARCHAR','Giữ số 0 đầu; không dùng làm phép tính'],
 ['Ngày / thời điểm','DATE / TIMESTAMPTZ','Phân biệt ngày nghiệp vụ và thời điểm'],
 ['Kỳ hóa đơn','DATE ngày đầu tháng','Kiểu ngày hợp lệ; CHECK ngày = 1']],
 [3.0,3.5,5.5],row_h=.65,size=18)

# 14
s=base('Khóa và ràng buộc toàn vẹn','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Cơ chế','Quy tắc được bảo vệ'],
 ['PK + FK','Mọi bảng có định danh; không tham chiếu bản ghi không tồn tại'],
 ['UNIQUE','Không trùng mã phòng trong khu, hóa đơn/kỳ hoặc mã giao dịch'],
 ['CHECK','Tiền hợp lệ; mới ≥ cũ; ngày hợp lệ; trạng thái đúng tập giá trị'],
 ['EXCLUDE','Không chồng khoảng thuê phòng hoặc khoảng cư trú của cùng thành viên'],
 ['Trigger + khóa hàng','Không thu vượt nợ; không sửa hóa đơn phát hành; kiểm tra liên bảng']],
 [3.0,9.0],row_h=.80,size=19)

# 15
s=base('Quan hệ và bảng trung gian','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Quan hệ','Ví dụ','Cách biểu diễn'],
 ['1–0..1','Tài khoản–hồ sơ; hợp đồng–quyết toán','UNIQUE trên FK; liên kết tùy chọn'],
 ['1–N','Khu–phòng; hợp đồng–hóa đơn; hóa đơn–giao dịch','FK ở phía nhiều'],
 ['N–N','Nhân viên–khu nhà','employee_assignments'],
 ['N–N theo thời gian','Người thuê–hợp đồng ở cùng','contract_members + khoảng cư trú'],
 ['N–N có lịch sử giá','Khu nhà–dịch vụ','service_price_configs + ngày hiệu lực']],
 [2.7,5.2,4.1],row_h=.80,size=18)

# 16
s=base('Chuẩn hóa dữ liệu','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Mức','Vấn đề cần tránh','Cách xử lý trong mô hình'],
 ['1NF','Nhiều ảnh trong một cột','maintenance_images: mỗi dòng một ảnh'],
 ['2NF','Thông tin chỉ phụ thuộc một phần khóa ghép','Không lưu họ tên ở thành viên; không lưu room_id ở chỉ số'],
 ['3NF','Thông tin danh mục lặp qua bảng quan hệ','Hồ sơ người và danh mục dịch vụ lưu tại thực thể riêng'],
 ['Dữ liệu suy ra','Nhiều cột tổng hợp có thể cập nhật lệch nhau','Tính lượng dùng, thành tiền, công nợ và trạng thái qua view']],
 [1.8,4.6,5.6],row_h=.94,size=18)

# 17
s=base('Phụ thuộc dữ liệu và bảo toàn lịch sử','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Phụ thuộc','Quyết định thiết kế'],
 ['tenant_id → thông tin cá nhân','Bảng thành viên chỉ tham chiếu hồ sơ người thuê'],
 ['contract_id → room_id → property_id','Chỉ số và sự cố truy ra khu/phòng qua hợp đồng'],
 ['Chỉ số mới − chỉ số cũ → lượng dùng','Tính lượng dùng; không nhập lại tại cảnh báo/hóa đơn'],
 ['Dòng hóa đơn + giao dịch → công nợ','Một view tính tổng, đã thu và số còn nợ']],
 [5.1,6.9],row_h=.90,size=18.5)
text(s,.8,6.48,11.9,.4,'Tên, đơn vị và giá snapshot là dữ kiện tại lúc lập hóa đơn; không đổi theo danh mục hiện tại.',17,True,ACCENT)

# 18
s=base('Đối chiếu nghiệp vụ và kiểm tra dữ liệu','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Nghiệp vụ','Dữ liệu chính','Tình huống đã kiểm tra'],
 ['Lập hợp đồng','rooms, contracts, contract_members','Chồng thời gian; phòng bảo trì; vượt sức chứa'],
 ['Lập/phát hành hóa đơn','meter_readings, invoices, invoice_items','Trùng kỳ; cảnh báo PENDING; sai chỉ số tham chiếu'],
 ['Thu tiền','invoices, payment_transactions','Tiền âm; trùng giao dịch; thu vượt công nợ'],
 ['Trả phòng','checkout_records và giao dịch cấn cọc','Snapshot trước cấn; không tính điện nước hai lần'],
 ['Gửi sự cố','maintenance_requests, maintenance_images','Sai người gửi; sai người đính kèm; ảnh rỗng']],
 [2.3,4.4,5.3],row_h=.78,size=17)
text(s,.8,6.62,11.8,.3,'PostgreSQL: 40 ca đạt; thu đồng thời 600 + 600 trên nợ 1.000 → nhận một lần, còn 400.',16,True,ACCENT)

# 19
s=base('Ví dụ hóa đơn và quyết toán','02  THIẾT KẾ CƠ SỞ DỮ LIỆU')
table(s,[['Khoản thu — ở đủ tháng 09/2026','Công thức','Thành tiền'],
 ['Tiền phòng','3.000.000 × 30/30','3.000.000 đ'],
 ['Điện','(1.120 − 1.000) × 3.500','420.000 đ'],
 ['Nước','(128 − 120) × 20.000','160.000 đ'],
 ['Internet','100.000 × 30/30','100.000 đ'],
 ['Tổng hóa đơn','','3.680.000 đ']],
 [5.0,4.3,2.7],y=1.7,row_h=.59,size=18)
text(s,.8,5.5,11.9,.4,'Đã thu 2.000.000 đ → còn nợ 1.680.000 đ.',22,True)
text(s,.8,6.1,11.9,.5,'Cọc 3.000.000 − thiệt hại 200.000 − nợ 1.680.000 = hoàn 1.120.000 đ.',19,False,ACCENT)
CONTENT[-1]['items']=['Dữ liệu minh họa; các khoản nợ khác giả định bằng 0.',
 'Cấn cọc 1.680.000 tạo DEPOSIT_OFFSET; hóa đơn hết nợ nhưng tiền mới nhận vẫn là 2.000.000.']

# 20
s=base('Phân công thực hiện','THÀNH VIÊN NHÓM')
table(s,[['Họ tên — MSSV','Phần việc thực hiện','Minh chứng','Đóng góp'],
 ['……………………','………………………………','Tài liệu / commit','…… %'],
 ['……………………','………………………………','Tài liệu / commit','…… %'],
 ['……………………','………………………………','Tài liệu / commit','…… %'],
 ['……………………','………………………………','Tài liệu / commit','…… %']],
 [3.2,4.4,2.6,1.8],row_h=.9,size=18)
text(s,.8,6.55,11.8,.3,'Tổng tỷ lệ đóng góp: 100%.',19,True)

# 21: supporting ERD, no presenter notes or instructions.
s=base('ERD — yêu cầu sửa chữa và ảnh','PHỤ LỤC')
a=entity(s,.65,2.0,'contracts',['PK  id','FK  primary_tenant_id'])
b=entity(s,4.95,2.0,'maintenance_requests',['PK  id','FK  contract_id','FK  submitted_by; assigned_to','title; priority; status'])
c=entity(s,9.25,2.0,'maintenance_images',['PK  id','FK  maintenance_request_id','FK  uploaded_by','url'])
d=entity(s,4.95,4.8,'users',['PK  id','role; status'])
relation(s,a,b);relation(s,b,c);relation(s,d,b,vertical=True)
route(s,[(8.05,5.5),(10.8,5.5),(10.8,c[1]+c[3])],'1 : N',(9.1,5.18,1.0))
text(s,.8,6.65,11.8,.3,'Không lưu lại khu/phòng ở sự cố; mỗi ảnh có bản ghi và người đính kèm riêng.',17,False,MUTED)
CONTENT[-1]['items']=['Sự cố gắn hợp đồng; khu/phòng truy ra qua hợp đồng.',
 'Người gửi, người được giao và người đính kèm đều tham chiếu users.',
 'Một yêu cầu có nhiều ảnh; mỗi dòng một URL.']

assert len(prs.slides)==21
prs.save(OUT)
with zipfile.ZipFile(OUT) as archive:
    assert not any(name.startswith('ppt/notesSlides/notesSlide') and name.endswith('.xml') for name in archive.namelist())
    assert archive.testzip() is None

# Matching source content contains no speaker notes or drafting guidance.
lines=['# Hệ thống quản lý nhà trọ','',
 'Phân tích yêu cầu và thiết kế cơ sở dữ liệu — học phần Kỹ năng lập trình nâng cao, N03.','',
 '**Phạm vi:** Mô hình mục tiêu 17 bảng; một chủ nhà, nhân viên theo khu và người đại diện hợp đồng.','',
 '**Nhóm:** ………………………………','']
for slide in CONTENT:
    lines+=['## Slide '+f'{slide["slide"]:02}'+ ' — '+slide['title'],'']
    for item in slide['items']:lines+=['- '+item]
    if slide['items']:lines+=['']
    if 'table' in slide:
        rows=slide['table'];lines+=['| '+' | '.join(str(x).replace('\n',' / ') for x in rows[0])+' |',
          '| '+' | '.join('---' for _ in rows[0])+' |']
        lines+=['| '+' | '.join(str(x).replace('\n',' / ') for x in row)+' |' for row in rows[1:]]
        lines+=['']

lines+=['# Phụ lục — danh mục Use Case','',
 '**A:** ADMIN. **S:** STAFF theo khu được giao. **T:** TENANT theo hợp đồng mình đại diện. Hồ sơ chưa gắn hợp đồng do ADMIN tiếp nhận; STAFF tra cứu hồ sơ có căn cứ trong khu được giao.','',
 '| Mã | Chức năng | Tác nhân | Dữ liệu chính |','| --- | --- | --- | --- |']
catalog=[
 ('Đăng nhập hệ thống','A/S/T','users'),('Đổi mật khẩu','A/S/T','users'),
 ('Tạo tài khoản','A','users'),('Cập nhật, khóa tài khoản','A','users'),('Phân công, thu hồi phân công nhân viên','A','employee_assignments'),
 ('Tạo khu nhà','A','properties'),('Cập nhật khu nhà','A','properties'),('Ngừng khai thác khu nhà','A','properties'),
 ('Tạo phòng','A','rooms'),('Cập nhật phòng','A/S','rooms'),('Chuyển phòng sang bảo trì','A/S','rooms'),('Xóa phòng chưa có phụ thuộc','A','rooms'),
 ('Tạo hồ sơ người thuê','A','tenant_profiles'),('Cập nhật hồ sơ người thuê','A','tenant_profiles'),('Tra cứu hồ sơ người thuê','A/S/T','tenant_profiles, contracts'),
 ('Lập hợp đồng thuê','A/S','contracts, rooms, tenant_profiles'),('Tra cứu hợp đồng','A/S/T','contracts'),('Cập nhật thành viên cư trú','A/S','contract_members, tenant_profiles'),
 ('Thiết lập dịch vụ và đơn giá','A','service_items, service_price_configs'),('Ghi chỉ số điện nước','A/S','meter_readings'),('Sửa chỉ số cuối chưa dùng trong hóa đơn','A/S','meter_readings'),
 ('Kiểm tra, xác nhận cảnh báo','A/S','consumption_alerts'),('Lập hóa đơn nháp','A/S','invoices, invoice_items'),('Xóa hóa đơn nháp','A/S','invoices, invoice_items'),
 ('Phát hành hóa đơn','A/S','invoices, consumption_alerts'),('Tra cứu hóa đơn đã phát hành','A/S/T','v_invoice_lines, v_invoice_balances'),
 ('Ghi nhận thanh toán','A/S','payment_transactions'),('Tra cứu công nợ','A/S/T','v_invoice_balances'),('Quyết toán trả phòng','A/S','contracts, checkout_records, payment_transactions'),
 ('Gửi yêu cầu sửa chữa','T','maintenance_requests'),('Cập nhật tiến độ sửa chữa','A/S','maintenance_requests'),('Theo dõi yêu cầu sửa chữa','A/S/T','maintenance_requests'),
 ('Tổng hợp tiền thu','A/S','payment_transactions, v_invoice_balances'),('Đính kèm ảnh sự cố','A/S/T','maintenance_images'),('Thiết lập ngưỡng cảnh báo theo khu','A','properties')]
for index,(name,actors,data) in enumerate(catalog,1):lines.append(f'| UC{index:02} | {name} | {actors} | {data} |')
lines+=['','## Quy tắc dữ liệu','',
 '- Một hồ sơ dùng chung cho các lần cư trú. Người đại diện không được thêm lần nữa như thành viên của chính hợp đồng.',
 '- Mã phòng duy nhất trong khu; không cho thuê chồng khoảng thời gian; không vượt sức chứa.',
 '- Ngày ra không tính tiền. Tiền phòng/phí tháng chia theo ngày hoặc ngày-người; điện nước dùng chênh chỉ số.',
 '- Kỳ là ngày đầu tháng. Một hợp đồng/kỳ có tối đa một hóa đơn; không phát hành khi còn cảnh báo PENDING.',
 '- Hóa đơn phát hành và chỉ số đã sử dụng không được sửa; giao dịch thu tiền chỉ được thêm mới.',
 '- Giữ tên, đơn vị, đơn giá tại lúc tính tiền. Tổng tiền và công nợ tính qua view, không lưu lặp.',
 '- Quyết toán giữ cọc và nợ snapshot trước cấn. Cấn cọc không phải dòng tiền thu mới; số quyết toán có thể âm.',
 '- Cùng một transaction phải giữ hợp đồng kết thúc, bản quyết toán và giao dịch cấn cọc nhất quán.','',
 '## Đặc tả Use Case chính','',
 '| Use Case | Tiền điều kiện | Luồng chính | Nhánh thay thế | Hậu điều kiện |',
 '| --- | --- | --- | --- | --- |',
 '| Phát hành hóa đơn | Đúng quyền; hóa đơn DRAFT | Kiểm tra ngày, giá, lượng tính phí và cảnh báo → ISSUED | Thiếu dữ liệu hoặc còn PENDING thì giữ nháp; sai chỉ số thì xóa nháp và lập lại | Snapshot và dữ liệu tham chiếu được bảo vệ |',
 '| Quyết toán trả phòng | Hợp đồng/ngày hợp lệ; chưa quyết toán | Hoàn tất mọi kỳ → chốt cọc/nợ/thiệt hại → cấn cọc → lưu kết quả | Thiếu hóa đơn thì bổ sung; dữ liệu phát hành không khớp thì dừng; cọc thiếu thì ghi phải thu thêm | Một bản quyết toán; công nợ và giao dịch cấn cọc đối chiếu được |','',
 '## Hồ sơ dữ liệu','',
 '- [Thiết kế và ràng buộc](DATABASE.md).',
 '- [Chứng minh chuẩn hóa](NORMALIZATION.md).',
 '- [DBML](quanlytro.dbml), [SQL](quanlytro.sql), [view và trigger](database-rules.sql).',
 '- [ERD PDF](diagrams/quanlytro-dbdiagram.pdf).',
 '- [Ca kiểm tra DB](database-cases.sql), [kết quả thực thi](database-validation.json).',
 '- [Luồng hóa đơn](diagrams/thuyet-trinh-luong-hoa-don.svg), [luồng trả phòng](diagrams/thuyet-trinh-luong-tra-phong.svg).','',
 '## Phiên bản triển khai','',
 'Đây là mô hình mục tiêu của bước 1–2. Backend 29 bảng V1–V9 là prototype riêng; chưa có migration/API chạy trên mô hình 17 bảng. Kết quả 40 ca và một ca đồng thời là kiểm tra SQL thiết kế trên database tạm, không phải kết quả kiểm thử API hoặc điểm chấm của giảng viên.','']
(DOCS/'NOI-DUNG-THUYET-TRINH-BUOC-1-2.md').write_text('\n'.join(lines),encoding='utf-8')
(DOCS/'presentation-content.json').write_text(json.dumps(CONTENT,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'Created {OUT.name}: {len(prs.slides)} slides, native text/tables/diagrams, no notes.')
